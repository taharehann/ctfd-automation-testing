import csv
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DOCS_DIR = os.path.join(BASE_DIR, "docs")
DATA_DIR = os.path.join(BASE_DIR, "data")
OUTPUT_FILE = os.path.join(DOCS_DIR, "CTFd_Test_Cases_Member2.xlsx")

# ---- Styles ----
FONT_FAMILY = "Segoe UI"
HEADER_FONT = Font(name=FONT_FAMILY, size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="2A4D69", end_color="2A4D69", fill_type="solid")
TITLE_FONT = Font(name=FONT_FAMILY, size=16, bold=True, color="1B365D")
SECTION_FONT = Font(name=FONT_FAMILY, size=12, bold=True, color="1B365D")
DATA_FONT = Font(name=FONT_FAMILY, size=10)
BOLD_FONT = Font(name=FONT_FAMILY, size=10, bold=True)
PASS_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="008000")
FAIL_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="C00000")

THIN_BORDER = Border(
    left=Side(style="thin", color="D3D3D3"),
    right=Side(style="thin", color="D3D3D3"),
    top=Side(style="thin", color="D3D3D3"),
    bottom=Side(style="thin", color="D3D3D3")
)
HEADER_BORDER = Border(
    left=Side(style="thin", color="FFFFFF"),
    right=Side(style="thin", color="FFFFFF"),
    top=Side(style="thin", color="1B365D"),
    bottom=Side(style="medium", color="1B365D")
)

# ---- Master Test Cases Data (40 test cases across Member 2's 3 modules) ----
TEST_CASES = [
    # Module A: Profile & Settings
    ("SET-STR1", "Profile & Settings", "notes/settings_test_design.md", "Structural / Baseline",
     "Settings page opens, heading is Settings, pre-populates username",
     "Logged-in user visits /settings", "Heading Settings, name input contains username", "-",
     "SettingsTests.test_settings_page_opens", "Pass"),
    ("SET-STR2", "Profile & Settings", "notes/settings_test_design.md", "Structural / Persistence",
     "Profile fields affiliation and country persist on page reload",
     "affiliation='Independent Security Researchers', country='US'", "Values persist in form after page reload", "-",
     "SettingsTests.test_profile_persistence_on_reload", "Pass"),
    ("SET-EC1", "Profile & Settings", "notes/settings_test_design.md", "Equivalence Class Partitioning",
     "Change user name to unique valid name",
     "name='user_{uid}'", "Your profile has been updated", "data/settings_data.csv",
     "SettingsTests.test_settings_data_cases_01", "Pass"),
    ("SET-EC3", "Profile & Settings", "notes/settings_test_design.md", "Equivalence Class Partitioning",
     "Change user name to existing name (admin)",
     "name='admin'", "User name has already been taken", "data/settings_data.csv",
     "SettingsTests.test_settings_data_cases_02", "Pass"),
    ("SET-EC4", "Profile & Settings", "notes/settings_test_design.md", "Equivalence Class Partitioning",
     "Change email with correct current password confirmation",
     "email='new_{uid}@test.com', confirm='Pass@123'", "Your profile has been updated", "data/settings_data.csv",
     "SettingsTests.test_settings_data_cases_03", "Pass"),
    ("SET-EC5", "Profile & Settings", "notes/settings_test_design.md", "Equivalence Class Partitioning",
     "Change email to existing email (admin@test.com)",
     "email='admin@test.com', confirm='Pass@123'", "Email address has already been used", "data/settings_data.csv",
     "SettingsTests.test_settings_data_cases_04", "Pass"),
    ("SET-EC6", "Profile & Settings", "notes/settings_test_design.md", "Boundary Value Analysis",
     "Change email without current password confirmation",
     "email='new_{uid}@test.com', confirm=''", "Please confirm your current password", "data/settings_data.csv",
     "SettingsTests.test_settings_data_cases_05", "Pass"),
    ("SET-EC7", "Profile & Settings", "notes/settings_test_design.md", "Equivalence Class Partitioning",
     "Change password with correct current password",
     "confirm='Pass@123', password='NewPass@456'", "Your profile has been updated", "data/settings_data.csv",
     "SettingsTests.test_settings_data_cases_06", "Pass"),
    ("SET-EC8", "Profile & Settings", "notes/settings_test_design.md", "Equivalence Class Partitioning",
     "Change password with incorrect current password",
     "confirm='Wrong@999', password='NewPass@456'", "Your previous password is incorrect", "data/settings_data.csv",
     "SettingsTests.test_settings_data_cases_07", "Pass"),
    ("SET-EC9", "Profile & Settings", "notes/settings_test_design.md", "Boundary Value Analysis",
     "Change password without current password confirmation",
     "confirm='', password='NewPass@456'", "Please confirm your current password", "data/settings_data.csv",
     "SettingsTests.test_settings_data_cases_08", "Pass"),
    ("SET-EC10", "Profile & Settings", "notes/settings_test_design.md", "Equivalence Class Partitioning",
     "Valid HTTPS website URL",
     "website='https://ctfd.io'", "Your profile has been updated", "data/settings_data.csv",
     "SettingsTests.test_settings_data_cases_09", "Pass"),
    ("SET-EC11", "Profile & Settings", "notes/settings_test_design.md", "Equivalence Class Partitioning",
     "Valid HTTP website URL",
     "website='http://example.com'", "Your profile has been updated", "data/settings_data.csv",
     "SettingsTests.test_settings_data_cases_10", "Pass"),
    ("SET-EC12", "Profile & Settings", "notes/settings_test_design.md", "Boundary Value Analysis",
     "Website URL missing protocol scheme",
     "website='example.com'", "Websites must be a proper URL / validation rejection", "data/settings_data.csv",
     "SettingsTests.test_settings_data_cases_11", "Pass"),
    ("SET-SYN1", "Profile & Settings", "notes/settings_test_design.md", "Syntax Testing",
     "Invalid URI protocol scheme (ftp://)",
     "website='ftp://ftp.example.com'", "Websites must be a proper URL / validation error", "data/settings_data.csv",
     "SettingsTests.test_settings_data_cases_12", "Pass"),
    ("SET-PW1", "Profile & Settings", "notes/settings_test_design.md", "Pairwise Testing",
     "Simultaneous combination of affiliation, website, and country",
     "affiliation='Security Lab', website='https://lab.org', country='US'", "Your profile has been updated", "data/settings_data.csv",
     "SettingsTests.test_settings_data_cases_13", "Pass"),

    # Module B: Users Directory & Public Profile
    ("USR-STR1", "Users Directory & Public Profile", "notes/users_test_design.md", "Structural / Baseline",
     "Users directory renders table columns and search form",
     "Visit /users", "User, Website, Affiliation, Country columns and search inputs", "-",
     "UsersTests.test_users_directory_page_structure", "Pass"),
    ("USR-UC1", "Users Directory & Public Profile", "notes/users_test_design.md", "Use Case / Scenario",
     "Navigation from directory to public profile",
     "Search user -> click user link", "Navigates to /users/<id>, heading matches username", "-",
     "UsersTests.test_directory_to_profile_navigation", "Pass"),
    ("USR-UC2", "Users Directory & Public Profile", "notes/users_test_design.md", "Use Case / Scenario",
     "Challenge solve reflection lifecycle on public profile",
     "Submit flag{test123} -> view /users/<id>", "Shows 100 points, Solves table with Test Challenge", "-",
     "UsersTests.test_user_profile_solve_reflection", "Pass"),
    ("USR-STR2", "Users Directory & Public Profile", "notes/users_test_design.md", "Structural / Integration",
     "Profile reflects affiliation, country flag, and website",
     "Set profile metadata in /settings -> visit /users/<id>", "Affiliation badge, country, and website link visible", "-",
     "UsersTests.test_user_profile_metadata_display", "Pass"),
    ("USR-EC1", "Users Directory & Public Profile", "notes/users_test_design.md", "Equivalence Class Partitioning",
     "Directory displays registered user accounts",
     "Visit /users", "Users table contains registered user rows", "data/users_data.csv",
     "UsersTests.test_users_data_driven_1", "Pass"),
    ("USR-EC3", "Users Directory & Public Profile", "notes/users_test_design.md", "Equivalence Class Partitioning",
     "Non-existent user profile returns 404",
     "Visit /users/999999", "404 Not Found error page", "data/users_data.csv",
     "UsersTests.test_users_data_driven_2", "Pass"),
    ("USR-EC4", "Users Directory & Public Profile", "notes/users_test_design.md", "Equivalence Class Partitioning",
     "User search matches registered user",
     "field=name, q={user_name}", "User row returned in results table", "data/users_data.csv",
     "UsersTests.test_users_data_driven_3", "Pass"),
    ("USR-EC5", "Users Directory & Public Profile", "notes/users_test_design.md", "Equivalence Class Partitioning",
     "User search matches no users",
     "field=name, q=nonexistent_user_xyz_999", "0 matching rows returned", "data/users_data.csv",
     "UsersTests.test_users_data_driven_4", "Pass"),
    ("USR-BVA1", "Users Directory & Public Profile", "notes/users_test_design.md", "Boundary Value Analysis",
     "Min score boundary: 0-point user listed in users directory",
     "Fresh user with 0 points", "User appears in /users table (unlike scoreboard)", "data/users_data.csv",
     "UsersTests.test_users_data_driven_5", "Pass"),
    ("USR-BVA3", "Users Directory & Public Profile", "notes/users_test_design.md", "Boundary Value Analysis",
     "Pagination min boundary: first page (page 1)",
     "/users?page=1", "First page loads with users table", "data/users_data.csv",
     "UsersTests.test_users_data_driven_6", "Pass"),
    ("USR-BVA4", "Users Directory & Public Profile", "notes/users_test_design.md", "Boundary Value Analysis",
     "Pagination out-of-bounds boundary: page 9999",
     "/users?page=9999", "Empty table, 0 rows, no crash", "data/users_data.csv",
     "UsersTests.test_users_data_driven_7", "Pass"),
    ("USR-EG1", "Users Directory & Public Profile", "notes/users_test_design.md", "Error Guessing",
     "Extremely large integer ID for profile",
     "/users/99999999", "404 Not Found without 500 server error", "data/users_data.csv",
     "UsersTests.test_users_data_driven_8", "Pass"),
    ("USR-EG2", "Users Directory & Public Profile", "notes/users_test_design.md", "Error Guessing",
     "Search query containing HTML script tags",
     "field=name, q=<script>", "Safely escaped, no execution, 0 results", "data/users_data.csv",
     "UsersTests.test_users_data_driven_9", "Pass"),

    # Module C: Admin Panel
    ("ADM-STR1", "Admin Panel", "notes/admin_test_design.md", "Structural / Baseline",
     "Admin dashboard loads with Pages, Users, Challenges sub-menus",
     "Login as admin -> /admin", "Title Admin Panel, navigation sub-menus present", "-",
     "AdminTests.test_admin_dashboard_navigation", "Pass"),
    ("ADM-UC1", "Admin Panel", "notes/admin_test_design.md", "Use Case / Scenario",
     "Custom page creation, navbar integration, and public rendering",
     "Create page rules -> visit /rules", "Listed in admin pages, visible in nav, content rendered", "-",
     "AdminTests.test_custom_page_publication_lifecycle", "Pass"),
    ("ADM-ST1", "Admin Panel", "notes/admin_test_design.md", "State Transition Testing",
     "User ban and unban lifecycle: Active -> Banned (403) -> Active",
     "Ban user in admin -> access /challenges -> Unban", "403 Forbidden while banned; normal access restored on unban", "-",
     "AdminTests.test_user_ban_and_unban_lifecycle", "Pass"),
    ("ADM-EC1", "Admin Panel", "notes/admin_test_design.md", "Cause-Effect Graphing",
     "Unauthenticated access to /admin redirected to login",
     "Unauthenticated visit /admin", "Redirected to /login", "data/admin_data.csv",
     "AdminTests.test_admin_data_driven_1", "Pass"),
    ("ADM-EC2", "Admin Panel", "notes/admin_test_design.md", "Cause-Effect Graphing",
     "Regular authenticated user access to /admin redirected to login",
     "Regular user visit /admin", "Redirected to /login", "data/admin_data.csv",
     "AdminTests.test_admin_data_driven_2", "Pass"),
    ("ADM-EC3", "Admin Panel", "notes/admin_test_design.md", "Cause-Effect Graphing",
     "Admin authenticated access to /admin loads dashboard",
     "Admin visit /admin", "Admin Panel dashboard loaded", "data/admin_data.csv",
     "AdminTests.test_admin_data_driven_3", "Pass"),
    ("ADM-EC4", "Admin Panel", "notes/admin_test_design.md", "Decision Table",
     "Published custom page is rendered publicly and displayed in navbar",
     "draft=False, hidden=False, auth_required=False", "Content rendered, link present in navbar", "data/admin_data.csv",
     "AdminTests.test_admin_data_driven_4", "Pass"),
    ("ADM-EC5", "Admin Panel", "notes/admin_test_design.md", "Decision Table",
     "Draft custom page returns 404 on public route",
     "draft=True, hidden=False, auth_required=False", "404 Not Found", "data/admin_data.csv",
     "AdminTests.test_admin_data_driven_5", "Pass"),
    ("ADM-EC6", "Admin Panel", "notes/admin_test_design.md", "Decision Table",
     "Auth-required custom page redirects unauthenticated visitors",
     "draft=False, hidden=False, auth_required=True", "Redirected to /login", "data/admin_data.csv",
     "AdminTests.test_admin_data_driven_6", "Pass"),
    ("ADM-BVA1", "Admin Panel", "notes/admin_test_design.md", "Boundary Value Analysis",
     "Min route length boundary (short route)",
     "Route: r{uid}", "Page created and published at /r{uid}", "data/admin_data.csv",
     "AdminTests.test_admin_data_driven_7", "Pass"),
    ("ADM-BVA2", "Admin Panel", "notes/admin_test_design.md", "Boundary Value Analysis",
     "Long hyphenated route format boundary",
     "Route: long-rules-guide-2026-{uid}", "Page created and published at /long-rules-guide-2026-{uid}", "data/admin_data.csv",
     "AdminTests.test_admin_data_driven_8", "Pass"),
    ("ADM-EG1", "Admin Panel", "notes/admin_test_design.md", "Error Guessing",
     "Non-existent admin route returns 404 gracefully",
     "/admin/nonexistent_xyz_999", "404 Not Found error page", "data/admin_data.csv",
     "AdminTests.test_admin_data_driven_9", "Pass"),
]


def auto_fit_columns(ws, min_width=12, max_width=50):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = str(cell.value or '')
            if cell.number_format and '%' in cell.number_format:
                val = "100.0%"
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, min_width), max_width)


def create_workbook():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # 1. Sheet: Test Cases
    ws_tc = wb.create_sheet(title="Test Cases")
    tc_headers = [
        "Test ID", "Module", "Design Ref (notes)", "Technique", "Description",
        "Input Data", "Expected Result", "Data Source", "Automated Test", "Status"
    ]
    ws_tc.append(tc_headers)
    for col_idx in range(1, len(tc_headers) + 1):
        cell = ws_tc.cell(1, col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = HEADER_BORDER
    ws_tc.row_dimensions[1].height = 28

    for row_idx, row_data in enumerate(TEST_CASES, start=2):
        ws_tc.append(row_data)
        for col_idx in range(1, len(row_data) + 1):
            cell = ws_tc.cell(row_idx, col_idx)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center")
            if col_idx in (1, 10):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if col_idx == 10:
                cell.font = PASS_FONT if cell.value == "Pass" else FAIL_FONT
        ws_tc.row_dimensions[row_idx].height = 20
    auto_fit_columns(ws_tc)

    # 2. Sheet: Settings Data
    ws_set = wb.create_sheet(title="Settings Data")
    with open(os.path.join(DATA_DIR, "settings_data.csv"), newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        for r_idx, row in enumerate(reader, start=1):
            ws_set.append(row)
            for c_idx in range(1, len(row) + 1):
                cell = ws_set.cell(r_idx, c_idx)
                cell.border = THIN_BORDER
                if r_idx == 1:
                    cell.font = HEADER_FONT
                    cell.fill = HEADER_FILL
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.font = DATA_FONT
                    cell.alignment = Alignment(vertical="center")
            ws_set.row_dimensions[r_idx].height = 22 if r_idx == 1 else 19
    auto_fit_columns(ws_set)

    # 3. Sheet: Users Data
    ws_usr = wb.create_sheet(title="Users Data")
    with open(os.path.join(DATA_DIR, "users_data.csv"), newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        for r_idx, row in enumerate(reader, start=1):
            ws_usr.append(row)
            for c_idx in range(1, len(row) + 1):
                cell = ws_usr.cell(r_idx, c_idx)
                cell.border = THIN_BORDER
                if r_idx == 1:
                    cell.font = HEADER_FONT
                    cell.fill = HEADER_FILL
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.font = DATA_FONT
                    cell.alignment = Alignment(vertical="center")
            ws_usr.row_dimensions[r_idx].height = 22 if r_idx == 1 else 19
    auto_fit_columns(ws_usr)

    # 4. Sheet: Admin Data
    ws_adm = wb.create_sheet(title="Admin Data")
    with open(os.path.join(DATA_DIR, "admin_data.csv"), newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        for r_idx, row in enumerate(reader, start=1):
            ws_adm.append(row)
            for c_idx in range(1, len(row) + 1):
                cell = ws_adm.cell(r_idx, c_idx)
                cell.border = THIN_BORDER
                if r_idx == 1:
                    cell.font = HEADER_FONT
                    cell.fill = HEADER_FILL
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.font = DATA_FONT
                    cell.alignment = Alignment(vertical="center")
            ws_adm.row_dimensions[r_idx].height = 22 if r_idx == 1 else 19
    auto_fit_columns(ws_adm)

    # 5. Sheet: Summary
    ws_sum = wb.create_sheet(title="Summary")
    ws_sum.cell(1, 1, "CTFd Test Execution Summary (Member 2)").font = TITLE_FONT
    ws_sum.row_dimensions[1].height = 30

    sum_headers = ["Module", "Total", "Pass", "Fail", "Not Run", "Pass Rate"]
    for c_idx, h in enumerate(sum_headers, start=1):
        cell = ws_sum.cell(3, c_idx, h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = HEADER_BORDER
    ws_sum.row_dimensions[3].height = 24

    modules = [
        "Profile & Settings",
        "Users Directory & Public Profile",
        "Admin Panel"
    ]
    max_tc_row = len(TEST_CASES) + 1  # 41

    for idx, mod in enumerate(modules, start=4):
        ws_sum.cell(idx, 1, mod).font = BOLD_FONT
        ws_sum.cell(idx, 2, f"=COUNTIF('Test Cases'!$B$2:$B${max_tc_row}, A{idx})")
        ws_sum.cell(idx, 3, f'=COUNTIFS(\'Test Cases\'!$B$2:$B${max_tc_row}, A{idx}, \'Test Cases\'!$J$2:$J${max_tc_row}, "Pass")')
        ws_sum.cell(idx, 4, f'=COUNTIFS(\'Test Cases\'!$B$2:$B${max_tc_row}, A{idx}, \'Test Cases\'!$J$2:$J${max_tc_row}, "Fail")')
        ws_sum.cell(idx, 5, f'=COUNTIFS(\'Test Cases\'!$B$2:$B${max_tc_row}, A{idx}, \'Test Cases\'!$J$2:$J${max_tc_row}, "Not Run")')
        ws_sum.cell(idx, 6, f'=IF(B{idx}=0, 0, C{idx}/B{idx})')
        ws_sum.cell(idx, 6).number_format = '0.0%'

        for c_idx in range(1, 7):
            cell = ws_sum.cell(idx, c_idx)
            cell.border = THIN_BORDER
            if c_idx > 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = DATA_FONT
        ws_sum.row_dimensions[idx].height = 20

    # Total row
    tot_row = len(modules) + 4  # 7
    ws_sum.cell(tot_row, 1, "Total").font = Font(name=FONT_FAMILY, size=11, bold=True)
    ws_sum.cell(tot_row, 2, f"=SUM(B4:B{tot_row-1})")
    ws_sum.cell(tot_row, 3, f"=SUM(C4:C{tot_row-1})")
    ws_sum.cell(tot_row, 4, f"=SUM(D4:D{tot_row-1})")
    ws_sum.cell(tot_row, 5, f"=SUM(E4:E{tot_row-1})")
    ws_sum.cell(tot_row, 6, f"=IF(B{tot_row}=0, 0, C{tot_row}/B{tot_row})")
    ws_sum.cell(tot_row, 6).number_format = '0.0%'

    for c_idx in range(1, 7):
        cell = ws_sum.cell(tot_row, c_idx)
        cell.font = Font(name=FONT_FAMILY, size=11, bold=True)
        cell.border = THIN_BORDER
        cell.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        if c_idx > 1:
            cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[tot_row].height = 22

    # Notes section
    ws_sum.cell(9, 1, "Notes & Verification Details").font = SECTION_FONT
    notes = [
        "Status values reflect automated Selenium test suite execution on local CTFd 3.8.8.",
        "Member 2 Modules: Module A (Profile & Settings), Module B (Users Directory & Profile), Module C (Admin Panel).",
        "Test Techniques Used: Equivalence Class Partitioning (ECP), Boundary Value Analysis (BVA), Syntax Testing, Pairwise Testing, Cause-Effect Graphing / Decision Table Testing, State Transition Testing, Use Case / Scenario Testing, Error Guessing.",
        "Change any value in the Status column of 'Test Cases' and this Summary table updates automatically via Excel formulas.",
        "Valid Status values: Pass, Fail, Not Run."
    ]
    for n_idx, note_text in enumerate(notes, start=10):
        cell = ws_sum.cell(n_idx, 1, f"• {note_text}")
        cell.font = DATA_FONT
        ws_sum.row_dimensions[n_idx].height = 18

    auto_fit_columns(ws_sum, min_width=14, max_width=45)
    ws_sum.column_dimensions["A"].width = 36

    wb.save(OUTPUT_FILE)
    print(f"Workbook saved successfully to: {OUTPUT_FILE}")

if __name__ == "__main__":
    create_workbook()
